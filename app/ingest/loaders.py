# -*- coding: utf-8 -*-
"""
文档解析器
==========

把不同格式的字节流统一转成纯文本，供检索层切分索引。
新增格式只需在此扩展，不影响其它层。
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Set

# 可直接按文本编码解码的格式
TEXT_SUFFIXES: Set[str] = {".txt", ".md", ".markdown"}
# 需要专门解析器的二进制格式
BINARY_SUFFIXES: Set[str] = {".pdf", ".docx", ".doc"}
SUPPORTED_SUFFIXES: Set[str] = TEXT_SUFFIXES | BINARY_SUFFIXES


class UnsupportedFormatError(ValueError):
    """格式不支持。"""


class ParseError(ValueError):
    """文件解析失败。"""


def _decode_text(data: bytes) -> str:
    """按常见中文编码依次尝试解码。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(data: bytes) -> str:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:                # pragma: no cover
        raise ParseError("未安装 PyPDF2，无法解析 PDF") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise ParseError(f"PDF 解析失败：{exc}") from exc


def _parse_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:                # pragma: no cover
        raise ParseError("未安装 python-docx，无法解析 Word 文档") from exc
    try:
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:         # 表格内容也别丢
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ParseError(f"Word 解析失败：{exc}") from exc


def extract_text(filename: str, data: bytes) -> str:
    """
    把上传文件解析为纯文本。

    抛出 UnsupportedFormatError / ParseError，由 API 层翻译成 400 响应。
    """
    suffix = Path(filename or "").suffix.lower()
    if suffix in TEXT_SUFFIXES:
        text = _decode_text(data)
    elif suffix == ".pdf":
        text = _parse_pdf(data)
    elif suffix in {".docx", ".doc"}:
        text = _parse_docx(data)
    else:
        raise UnsupportedFormatError(f"不支持的文件类型：{suffix or '未知'}")

    if not text.strip():
        raise ParseError("文件中没有可提取的文本内容")
    return text
